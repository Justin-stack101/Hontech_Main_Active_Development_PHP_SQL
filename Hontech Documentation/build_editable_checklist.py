import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import shutil
import zipfile
import os

def create_revised_workbook():
    template_path = "2025 BLANK RO UPDATED.xlsx"
    backup_path = "2025 BLANK RO UPDATED_original_backup.xlsx"
    revised_template = "2025 BLANK RO UPDATED.xlsx"
    ht_jo_output = "2025_RO_HT-JO-8268.xlsx"

    if not os.path.exists(backup_path):
        shutil.copy2(template_path, backup_path)
        print(f"Backed up original template to {backup_path}")

    wb = openpyxl.load_workbook(template_path)
    ws = wb['CHECKLIST']

    # Set column widths
    col_widths = {
        'A': 34.0,  # Left Item Description
        'B': 4.5,   # Left Status 🟩
        'C': 4.5,   # Left Status 🟨
        'D': 4.5,   # Left Status 🟥
        'E': 2.0,   # Spacer between left & right
        'F': 12.0,  # Right Component / Pos
        'G': 14.0,  # Right Wear Pattern / Input
        'H': 14.0,  # Right Tread / Thickness
        'I': 4.5,   # Right Status 🟩
        'J': 4.5,   # Right Status 🟨
        'K': 4.5,   # Right Status 🟥
        'L': 9.0,   # Date Label
        'M': 13.0   # Date Value
    }
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    # Define common styles
    font_family = "Segoe UI"
    f_title = Font(name=font_family, size=10, bold=True, color="FFFFFF")
    f_header = Font(name=font_family, size=9, bold=True)
    f_item = Font(name=font_family, size=8)
    f_item_bold = Font(name=font_family, size=8, bold=True)
    f_small = Font(name=font_family, size=7, italic=True)
    f_mark = Font(name=font_family, size=9, bold=True)
    
    fill_header_gray = PatternFill(start_color="808080", end_color="808080", fill_type="solid")
    fill_sub_gray = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    fill_green = PatternFill(start_color="00B050", end_color="00B050", fill_type="solid")
    fill_yellow = PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid")
    fill_red = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
    fill_light_green = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    fill_light_yellow = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    fill_light_red = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")

    thin_border_side = Side(style='thin', color='A0A0A0')
    thick_border_side = Side(style='medium', color='404040')
    
    box_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    bottom_line = Border(bottom=thin_border_side)

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center", wrap_text=True)

    def style_range(ws, cell_range, font=None, fill=None, border=None, alignment=None):
        for row in ws[cell_range]:
            for cell in row:
                if font: cell.font = font
                if fill: cell.fill = fill
                if border: cell.border = border
                if alignment: cell.alignment = alignment

    # Row 7: Fuel Level row next to VEHICLE YEAR MODEL
    ws['H7'] = "FUEL LEVEL:"
    ws['H7'].font = f_header
    ws['H7'].alignment = align_right
    fuel_levels = [('I7', 'E'), ('J7', '1/4'), ('K7', '1/2'), ('L7', '3/4'), ('M7', 'F')]
    for cell_id, val in fuel_levels:
        ws[cell_id] = val
        ws[cell_id].font = f_header
        ws[cell_id].alignment = align_center
        ws[cell_id].border = box_border

    # Row 9: Legend Banner
    # Left Legend
    ws.merge_cells('B9:C9')
    ws['B9'] = "Satisfactory"
    ws['B9'].font = Font(name=font_family, size=8, bold=True, color="FFFFFF")
    ws['B9'].fill = fill_green
    ws['B9'].alignment = align_center
    ws['B9'].border = box_border
    ws['C9'].border = box_border

    ws.merge_cells('D9:G9')
    ws['D9'] = "May Require Future Attention"
    ws['D9'].font = Font(name=font_family, size=8, bold=True)
    ws['D9'].fill = fill_yellow
    ws['D9'].alignment = align_center
    ws['D9'].border = box_border
    for c in ['E9', 'F9', 'G9']: ws[c].border = box_border

    ws.merge_cells('H9:K9')
    ws['H9'] = "Requires Immediate Attention"
    ws['H9'].font = Font(name=font_family, size=8, bold=True, color="FFFFFF")
    ws['H9'].fill = fill_red
    ws['H9'].alignment = align_center
    ws['H9'].border = box_border
    for c in ['I9', 'J9', 'K9']: ws[c].border = box_border

    # Row 11: Section Headers (Left: Interior/Exterior, Right: Tire Condition)
    ws['A11'] = "Interior / Exterior"
    ws['A11'].font = f_title
    ws['A11'].fill = fill_header_gray
    ws['A11'].alignment = align_left
    ws['A11'].border = box_border
    
    for c, fill_color in [('B11', fill_green), ('C11', fill_yellow), ('D11', fill_red)]:
        ws[c].fill = fill_color
        ws[c].border = box_border

    ws.merge_cells('F11:H11')
    ws['F11'] = "Tire Condition"
    ws['F11'].font = f_title
    ws['F11'].fill = fill_header_gray
    ws['F11'].alignment = align_center
    ws['F11'].border = box_border
    ws['G11'].border = box_border
    ws['H11'].border = box_border

    for c, fill_color in [('I11', fill_green), ('J11', fill_yellow), ('K11', fill_red)]:
        ws[c].fill = fill_color
        ws[c].border = box_border

    # Left Rows 12-18 (Interior/Exterior items)
    left_interior = [
        "Headlights (high/low)/Taillights/Brake lights/Hazards/Signals",
        "Interior light",
        "Windshield washer spray/Wiper operation/Blades/Condition",
        "Parking brake",
        "Horn operation",
        "Clutch operation (if applicable)",
        "Micron cabin filter**"
    ]
    for idx, item in enumerate(left_interior):
        r = 12 + idx
        ws[f'A{r}'] = item
        ws[f'A{r}'].font = f_item
        ws[f'A{r}'].alignment = align_left
        ws[f'A{r}'].border = box_border
        for col_l in ['B', 'C', 'D']:
            ws[f'{col_l}{r}'].border = box_border
            ws[f'{col_l}{r}'].alignment = align_center

    # Right Rows 12-16 (Tire condition items)
    tire_items = [
        ("Left Front", "Wear pattern: ____", "Tire tread: ___ 32nds"),
        ("Right Front", "Wear pattern: ____", "Tire tread: ___ 32nds"),
        ("Left Rear", "Wear pattern: ____", "Tire tread: ___ 32nds"),
        ("Right Rear", "Wear pattern: ____", "Tire tread: ___ 32nds"),
        ("Spare", "Wear pattern: ____", "Tire tread: ___ 32nds")
    ]
    for idx, (pos, wear, tread) in enumerate(tire_items):
        r = 12 + idx
        ws[f'F{r}'] = pos
        ws[f'F{r}'].font = f_item_bold
        ws[f'F{r}'].border = box_border
        ws[f'F{r}'].alignment = align_left

        ws[f'G{r}'] = wear
        ws[f'G{r}'].font = f_item
        ws[f'G{r}'].border = box_border
        ws[f'G{r}'].alignment = align_left

        ws[f'H{r}'] = tread
        ws[f'H{r}'].font = f_item
        ws[f'H{r}'].border = box_border
        ws[f'H{r}'].alignment = align_left

        for col_r in ['I', 'J', 'K']:
            ws[f'{col_r}{r}'].border = box_border
            ws[f'{col_r}{r}'].alignment = align_center

    # Row 17 & 18 Right: Tire Inflation
    ws.merge_cells('F17:H17')
    ws['F17'] = "Front tire inflation set to: _______ psi"
    ws['F17'].font = f_item
    ws['F17'].border = box_border
    for c in ['G17', 'H17']: ws[c].border = box_border

    ws.merge_cells('F18:H18')
    ws['F18'] = "Rear tire inflation set to: _______ psi"
    ws['F18'].font = f_item
    ws['F18'].border = box_border
    for c in ['G18', 'H18']: ws[c].border = box_border

    # Row 20: Section Headers (Left: Battery Performance, Right: Brake Condition)
    ws['A20'] = "Battery Performance"
    ws['A20'].font = f_title
    ws['A20'].fill = fill_header_gray
    ws['A20'].alignment = align_left
    ws['A20'].border = box_border
    for c, fill_color in [('B20', fill_green), ('C20', fill_yellow), ('D20', fill_red)]:
        ws[c].fill = fill_color
        ws[c].border = box_border

    ws.merge_cells('F20:H20')
    ws['F20'] = "Brake Condition"
    ws['F20'].font = f_title
    ws['F20'].fill = fill_header_gray
    ws['F20'].alignment = align_center
    ws['F20'].border = box_border
    ws['G20'].border = box_border
    ws['H20'].border = box_border
    for c, fill_color in [('I20', fill_green), ('J20', fill_yellow), ('K20', fill_red)]:
        ws[c].fill = fill_color
        ws[c].border = box_border

    # Left Rows 21-22: Battery items
    ws['A21'] = "Good"
    ws['A21'].font = f_item
    ws['A21'].border = box_border
    for c in ['B21', 'C21', 'D21']: ws[c].border = box_border; ws[c].alignment = align_center

    ws['A22'] = "Replace"
    ws['A22'].font = f_item
    ws['A22'].border = box_border
    for c in ['B22', 'C22', 'D22']: ws[c].border = box_border; ws[c].alignment = align_center

    # Right Rows 21-24: Brake items
    brake_items = [
        ("Left Front", "Thickness: _____ mms"),
        ("Right Front", "Thickness: _____ mms"),
        ("Left Rear", "Thickness: _____ mms"),
        ("Right Rear", "Thickness: _____ mms")
    ]
    for idx, (pos, thick) in enumerate(brake_items):
        r = 21 + idx
        ws[f'F{r}'] = pos
        ws[f'F{r}'].font = f_item_bold
        ws[f'F{r}'].border = box_border
        ws[f'F{r}'].alignment = align_left

        ws.merge_cells(f'G{r}:H{r}')
        ws[f'G{r}'] = thick
        ws[f'G{r}'].font = f_item
        ws[f'G{r}'].border = box_border
        ws[f'H{r}'].border = box_border
        ws[f'G{r}'].alignment = align_left

        for col_r in ['I', 'J', 'K']:
            ws[f'{col_r}{r}'].border = box_border
            ws[f'{col_r}{r}'].alignment = align_center

    # Row 25 Right: Brakes not inspected checkbox
    ws.merge_cells('F25:H25')
    ws['F25'] = "Brakes not inspected on this visit  [   ]"
    ws['F25'].font = f_item
    ws['F25'].border = box_border
    for c in ['G25', 'H25']: ws[c].border = box_border

    # Row 24 Left: Under Hood Header
    ws['A24'] = "Under Hood"
    ws['A24'].font = f_title
    ws['A24'].fill = fill_header_gray
    ws['A24'].alignment = align_left
    ws['A24'].border = box_border
    for c, fill_color in [('B24', fill_green), ('C24', fill_yellow), ('D24', fill_red)]:
        ws[c].fill = fill_color
        ws[c].border = box_border

    under_hood_items = [
        "Fluid levels: Oil/Coolant/Power steering/Brake/Washer/ATF",
        "Air filter condition**",
        "External drive belts and radiator hoses",
        "Hydraulic clutch reservoir fluid (M/T vehicles)"
    ]
    for idx, item in enumerate(under_hood_items):
        r = 25 + idx
        ws[f'A{r}'] = item
        ws[f'A{r}'].font = f_item
        ws[f'A{r}'].border = box_border
        ws[f'A{r}'].alignment = align_left
        for col_l in ['B', 'C', 'D']:
            ws[f'{col_l}{r}'].border = box_border
            ws[f'{col_l}{r}'].alignment = align_center

    # Row 27 Right: Damage Banner
    ws.merge_cells('F27:K27')
    ws['F27'] = "Please Indicate Areas of External Damage or Wear"
    ws['F27'].font = Font(name=font_family, size=8, bold=True, italic=True)
    ws['F27'].fill = fill_sub_gray
    ws['F27'].alignment = align_left
    for c in ['F27', 'G27', 'H27', 'I27', 'J27', 'K27']: ws[c].border = box_border

    # Damage area box borders (Rows 28 to 44)
    for r in range(28, 45):
        for c in ['F', 'G', 'H', 'I', 'J', 'K']:
            ws[f'{c}{r}'].border = Border(
                left=thin_border_side if c == 'F' else None,
                right=thin_border_side if c == 'K' else None,
                top=thin_border_side if r == 28 else None,
                bottom=thin_border_side if r == 44 else None
            )

    # Row 30 Left: Under Vehicle Header
    ws['A30'] = "Under Vehicle"
    ws['A30'].font = f_title
    ws['A30'].fill = fill_header_gray
    ws['A30'].alignment = align_left
    ws['A30'].border = box_border
    for c, fill_color in [('B30', fill_green), ('C30', fill_yellow), ('D30', fill_red)]:
        ws[c].fill = fill_color
        ws[c].border = box_border

    under_vehicle_items = [
        "Brake lines / Hoses / Parking brake cable",
        "Shock absorbers / Struts / Suspension / Tie rod ends & boots",
        "Exhaust system",
        "Engine oil and/or fluid leaks",
        "Drive shaft boots / Constant velocity boots and bands"
    ]
    for idx, item in enumerate(under_vehicle_items):
        r = 31 + idx
        ws[f'A{r}'] = item
        ws[f'A{r}'].font = f_item
        ws[f'A{r}'].border = box_border
        ws[f'A{r}'].alignment = align_left
        for col_l in ['B', 'C', 'D']:
            ws[f'{col_l}{r}'].border = box_border
            ws[f'{col_l}{r}'].alignment = align_center

    # Row 37 Left: Comments Header
    ws['A37'] = "Comments"
    ws['A37'].font = f_title
    ws['A37'].fill = fill_header_gray
    ws['A37'].alignment = align_left
    ws['A37'].border = box_border
    for c in ['B37', 'C37', 'D37']:
        ws[c].fill = fill_header_gray
        ws[c].border = box_border

    # Comments lines (Rows 38 to 44)
    for r in range(38, 45):
        ws.merge_cells(f'A{r}:D{r}')
        ws[f'A{r}'].font = f_item
        ws[f'A{r}'].alignment = align_left
        for col_l in ['A', 'B', 'C', 'D']:
            ws[f'{col_l}{r}'].border = Border(
                left=thin_border_side if col_l == 'A' else None,
                right=thin_border_side if col_l == 'D' else None,
                bottom=thin_border_side
            )

    # Footnotes (Row 46)
    ws['A46'] = "*Note: Brake fluid NOT filled - fluid level indicates pad wear"
    ws['A46'].font = f_small
    ws['F46'] = "**Refer to maintenance schedule"
    ws['F46'].font = f_small

    # Save revised blank template
    wb.save(revised_template)
    print(f"Saved revised editable template to {revised_template}")

    # Now create HT-JO-8268 populated version
    wb_ht = openpyxl.load_workbook(revised_template)
    jo_sheet = wb_ht['J.O.']
    checklist_sheet = wb_ht['CHECKLIST']

    # Populate J.O. fields
    jo_sheet['C10'] = "Juan Dela Cruz"
    jo_sheet['K5'] = "2026-09-19"
    jo_sheet['K10'] = "ABC 1234"
    jo_sheet['H10'] = "N/A"

    # Populate CHECKLIST specific fields
    checklist_sheet['K7'] = "✓ 1/2" # Fuel level 1/2 from PDF
    checklist_sheet['K7'].fill = fill_light_yellow

    # Technician name
    checklist_sheet['B62'] = "TECHNICIAN NAME: Manney Sarol"
    checklist_sheet['B62'].font = f_header

    # Comments
    checklist_sheet['A38'] = "Vehicle intake inspection cleared. No critical defects noted."
    checklist_sheet['A38'].font = f_item_bold

    # Mark Satisfactory (Green check) on items
    # Interior / exterior
    for r in range(12, 19):
        checklist_sheet[f'B{r}'] = "✓"
        checklist_sheet[f'B{r}'].font = f_mark
        checklist_sheet[f'B{r}'].fill = fill_light_green

    # Battery
    checklist_sheet['B21'] = "✓"
    checklist_sheet['B21'].font = f_mark
    checklist_sheet['B21'].fill = fill_light_green

    # Under hood
    for r in range(25, 29):
        checklist_sheet[f'B{r}'] = "✓"
        checklist_sheet[f'B{r}'].font = f_mark
        checklist_sheet[f'B{r}'].fill = fill_light_green

    # Under vehicle
    for r in range(31, 36):
        checklist_sheet[f'B{r}'] = "✓"
        checklist_sheet[f'B{r}'].font = f_mark
        checklist_sheet[f'B{r}'].fill = fill_light_green

    # Tires
    for r in range(12, 17):
        checklist_sheet[f'I{r}'] = "✓"
        checklist_sheet[f'I{r}'].font = f_mark
        checklist_sheet[f'I{r}'].fill = fill_light_green

    # Brakes
    for r in range(21, 25):
        checklist_sheet[f'I{r}'] = "✓"
        checklist_sheet[f'I{r}'].font = f_mark
        checklist_sheet[f'I{r}'].fill = fill_light_green

    wb_ht.save(ht_jo_output)
    print(f"Saved populated HT-JO-8268 workbook to {ht_jo_output}")

create_revised_workbook()
